from io import BytesIO

import pytest
from docx import Document

from app.profile.service import (
    MAX_CV_BYTES,
    CvValidationError,
    extract_cv_text,
    save_candidate_profile,
)


class _Result:
    def first(self):
        return None


class _Session:
    def exec(self, _query):
        return _Result()

    def add(self, profile) -> None:
        self.profile = profile

    def commit(self) -> None:
        pass

    def refresh(self, profile) -> None:
        profile.id = 1


def test_extracts_utf8_text_cv() -> None:
    text, extension = extract_cv_text("özgeçmiş.TXT", b"Python\n\n.NET")

    assert text == "Python\n.NET"
    assert extension == ".txt"


def test_extracts_docx_cv() -> None:
    output = BytesIO()
    document = Document()
    document.add_paragraph("Onur Başkan")
    document.add_paragraph("Backend Developer")
    document.save(output)

    text, extension = extract_cv_text("cv.docx", output.getvalue())

    assert text == "Onur Başkan\nBackend Developer"
    assert extension == ".docx"


@pytest.mark.parametrize("filename", ["cv.exe", "cv.doc", "cv.jpg"])
def test_rejects_unsupported_cv_types(filename: str) -> None:
    with pytest.raises(CvValidationError, match="PDF, DOCX veya TXT"):
        extract_cv_text(filename, b"content")


def test_rejects_oversized_cv() -> None:
    with pytest.raises(CvValidationError, match="en fazla 5 MB"):
        extract_cv_text("cv.txt", b"x" * (MAX_CV_BYTES + 1))


def test_rejects_empty_extracted_text() -> None:
    with pytest.raises(CvValidationError, match="okunabilir metin"):
        extract_cv_text("cv.txt", b"  \n")


def test_saves_normalized_profile_and_local_file(tmp_path) -> None:
    session = _Session()

    profile = save_candidate_profile(
        session,  # type: ignore[arg-type]
        name="  Onur Başkan  ",
        filename="onur.txt",
        content=b"Python\n\nFastAPI",
        upload_dir=tmp_path,
    )

    assert profile.id == 1
    assert profile.name == "Onur Başkan"
    assert profile.cv_text == "Python\nFastAPI"
    assert (tmp_path / "candidate-cv.txt").read_bytes() == b"Python\n\nFastAPI"


def test_rejects_blank_candidate_name(tmp_path) -> None:
    with pytest.raises(CvValidationError, match="en az 2 karakter"):
        save_candidate_profile(
            _Session(),  # type: ignore[arg-type]
            name="  ",
            filename="cv.txt",
            content=b"Python",
            upload_dir=tmp_path,
        )
